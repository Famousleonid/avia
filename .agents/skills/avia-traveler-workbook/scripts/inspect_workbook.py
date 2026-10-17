"""Read-only XLSX inventory with cached/formula values, cell provenance and styles."""
import argparse
import hashlib
import json
from pathlib import Path
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('source')
parser.add_argument('output')
args = parser.parse_args()
source = Path(args.source)
book = openpyxl.load_workbook(source, data_only=False)
cached = openpyxl.load_workbook(source, data_only=True)
result = {'source': str(source), 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
          'color_theme_xml': book.loaded_theme.decode('utf-8') if book.loaded_theme else None,
          'indexed_colors': list(book._colors), 'sheets': []}
for sheet in book:
    rows = []
    for row in sheet:
        cells = []
        for cell in row:
            if cell.value is None:
                continue
            cells.append({'cell': cell.coordinate, 'value': cell.value, 'cached': cached[sheet.title][cell.coordinate].value,
                          'type': cell.data_type, 'strike': bool(cell.font.strike), 'bold': bool(cell.font.bold),
                          'hidden_row': bool(sheet.row_dimensions[cell.row].hidden)})
        if cells:
            rows.append(cells)
    color = sheet.sheet_properties.tabColor
    tab_color = dict(color) if color is not None else None
    result['sheets'].append({'name': sheet.title, 'state': sheet.sheet_state, 'tab_color': tab_color,
        'rows': sheet.max_row, 'cols': sheet.max_column, 'merges': [str(x) for x in sheet.merged_cells.ranges], 'cells': rows})
Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
for sheet in result['sheets']:
    print(f"{sheet['name']} | {sheet['state']} | tab_color={sheet['tab_color']} | {sheet['rows']}x{sheet['cols']} | {sum(len(row) for row in sheet['cells'])} cells")
